#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
净值数据更新脚本
读取 Excel 净值文件，通过统一调度层对齐沪深300基准，生成各产品 JSON 数据
遵循金融数据源优先级: THS > TuShare Pro > AKShare
"""

import os
import sys
import json
import logging
from datetime import datetime
from pathlib import Path

# 项目根目录
PROJECT_ROOT = Path(__file__).parent.parent
DATA_DIR = PROJECT_ROOT / 'data'
sys.path.insert(0, str(Path(__file__).parent))

try:
    from financial_data_provider import provider
except ImportError:
    provider = None

# 产品配置
PRODUCTS = {
    'zhouqi': {
        'name': '泽鑫周期',
        'pattern': 'zhouqi_*.xlsx',
        'json_path': PROJECT_ROOT / 'zhouqi' / 'combined-data.json',
    },
    'jiazhi': {
        'name': '泽鑫价值',
        'pattern': 'jiazhi_*.xlsx',
        'json_path': PROJECT_ROOT / 'jiazhi' / 'combined-data.json',
    }
}


def get_latest_excel(pattern):
    """获取最新的 Excel 文件"""
    files = list(DATA_DIR.glob(pattern))
    if not files:
        # 兼容 .xls
        alt_pattern = pattern.replace('.xlsx', '.xls')
        files = list(DATA_DIR.glob(alt_pattern))
    if not files:
        return None
    return max(files, key=lambda f: f.stat().st_mtime)


def read_excel_nav(filepath):
    """读取 Excel 净值数据，自动兼容 .xls 和 .xlsx，并支持多列净值提取累计净值"""
    filepath = Path(filepath)
    records = []

    if filepath.suffix.lower() == '.xlsx':
        import openpyxl
        wb = openpyxl.load_workbook(str(filepath), data_only=True)
        sheet = wb.active

        # 寻找表头行
        header_row = -1
        date_col = 0
        nav_col = 1
        cum_nav_col = -1

        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if r_idx > 15:
                break
            row_strs = [str(c).strip() if c is not None else '' for c in row]
            for c_idx, val in enumerate(row_strs):
                if '净值日期' in val or '日期' in val:
                    header_row = r_idx
                    date_col = c_idx
                if '累计单位净值' in val or '累计净值' in val:
                    cum_nav_col = c_idx
                elif '单位净值' in val and cum_nav_col == -1:
                    nav_col = c_idx

        # 若有累计单位净值列，优先作为累计净值
        target_nav_col = cum_nav_col if cum_nav_col != -1 else nav_col

        for r_idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if header_row != -1 and r_idx <= header_row:
                continue
            if not row or len(row) <= target_nav_col:
                continue
            c_date = row[date_col]
            c_nav = row[target_nav_col]

            if c_date is None or c_nav is None:
                continue

            # 处理日期
            dt_str = None
            if isinstance(c_date, datetime):
                dt_str = c_date.strftime('%Y-%m-%d')
            else:
                s = str(c_date).strip()
                for fmt in ('%Y-%m-%d', '%Y/%m/%d', '%Y%m%d'):
                    try:
                        dt_str = datetime.strptime(s, fmt).strftime('%Y-%m-%d')
                        break
                    except ValueError:
                        pass
            if not dt_str:
                continue

            try:
                nav_val = float(c_nav)
                records.append((dt_str, nav_val))
            except (ValueError, TypeError):
                continue

    else:
        import xlrd
        wb = xlrd.open_workbook(str(filepath))
        sh = wb.sheet_by_index(0)

        for row in range(sh.nrows):
            cell_date = sh.cell_value(row, 0)
            cell_nav = sh.cell_value(row, 1)

            if isinstance(cell_date, float):
                date = xlrd.xldate_as_datetime(cell_date, wb.datemode)
                dt_str = date.strftime('%Y-%m-%d')
            else:
                try:
                    dt_str = datetime.strptime(str(cell_date).strip(), '%Y-%m-%d').strftime('%Y-%m-%d')
                except ValueError:
                    continue

            try:
                nav = float(cell_nav)
                records.append((dt_str, nav))
            except (ValueError, TypeError):
                continue

    records.sort(key=lambda x: x[0])
    return records


def get_benchmark_data(start_date):
    """通过统一金融调度层获取沪深300数据 (THS > TuShare > AKShare)"""
    if provider:
        return provider.get_hs300_daily(start_date)
    else:
        import akshare as ak
        df = ak.stock_zh_index_daily(symbol="sh000300")
        df['date'] = df['date'].astype(str)
        df = df[df['date'] >= start_date].sort_values('date')
        return dict(zip(df['date'], df['close'].astype(float)))


def align_data(records, benchmark_dict):
    """对齐基金数据和基准数据，基准归一化到起始日=1.0"""
    fund_start = records[0][0]
    benchmark_filtered = {d: v for d, v in benchmark_dict.items() if d >= fund_start}

    if not benchmark_filtered:
        print("  基准数据为空")
        return None

    bench_dates_sorted = sorted(benchmark_filtered.keys())
    bench_start_val = benchmark_filtered[bench_dates_sorted[0]]

    aligned_dates = []
    aligned_nav = []
    aligned_benchmark = []

    for date, nav in records:
        if date in benchmark_filtered:
            aligned_dates.append(date)
            aligned_nav.append(round(nav, 4))
            aligned_benchmark.append(round(benchmark_filtered[date] / bench_start_val, 4))

    print(f"  对齐数据: {len(aligned_dates)} 条 ({aligned_dates[0]} ~ {aligned_dates[-1]})")
    return {
        'dates': aligned_dates,
        'nav': aligned_nav,
        'benchmark': aligned_benchmark
    }


def process_clients():
    """处理客户信息表，生成 clients.json"""
    client_files = list(DATA_DIR.glob('客户信息*.xls')) + list(DATA_DIR.glob('客户信息*.xlsx'))
    if not client_files:
        print("  未找到客户信息表，跳过")
        return False

    client_file = sorted(client_files)[-1]
    print(f"  处理客户信息表: {client_file.name}")

    try:
        import xlrd
        import re
        wb = xlrd.open_workbook(str(client_file))
        sh = wb.sheet_by_index(0)

        product_code_map = {
            '晟孚泽鑫周期私募证券投资基金': 'zhouqi',
            '晟孚泽鑫价值私募证券投资基金': 'jiazhi',
            '周期': 'zhouqi',
            '价值': 'jiazhi'
        }

        clients = {}
        for r in range(1, sh.nrows):
            row = [sh.cell_value(r, c) for c in range(sh.ncols)]
            name = str(row[1]).strip()
            phone_raw = str(int(row[2]) if isinstance(row[2], float) else row[2]).strip()
            last4 = phone_raw[-4:]
            p_field = str(row[3]).strip()
            date_val = str(row[4]).strip()

            p_names = [p.strip() for p in re.split(r'[；;]+', p_field) if p.strip()]
            products = []
            for p_name in p_names:
                code = product_code_map.get(p_name, 'zhouqi' if '周期' in p_name else 'jiazhi')
                products.append({'name': p_name, 'code': code})

            clients[last4] = {
                'name': name,
                'phone': phone_raw,
                'products': products,
                'date': date_val
            }

        # 保留专业机构账号
        clients['admin'] = {
            'name': '专业机构',
            'password': 'admin',
            'role': 'institution',
            'products': [
                {'name': '晟孚泽鑫周期私募证券投资基金', 'code': 'zhouqi'},
                {'name': '晟孚泽鑫价值私募证券投资基金', 'code': 'jiazhi'}
            ],
            'date': '2026.09.10'
        }

        json_path = PROJECT_ROOT / 'clients.json'
        with open(str(json_path), 'w', encoding='utf-8') as f:
            json.dump(clients, f, ensure_ascii=False, indent=2)

        print(f"  客户信息已更新: {json_path}")
        return True
    except Exception as e:
        print(f"  处理客户信息失败: {e}")
        return False


def main():
    print(f"工作目录: {PROJECT_ROOT}")
    benchmark_dict = get_benchmark_data('2024-07-01')

    for key, config in PRODUCTS.items():
        print(f"\n{'='*50}")
        print(f"处理产品: {config['name']} ({key})")
        print(f"{'='*50}")

        excel_path = get_latest_excel(config['pattern'])
        if not excel_path:
            print(f"  未找到 {config['pattern']} 文件，跳过")
            continue

        print(f"  读取: {excel_path.name}")
        records = read_excel_nav(excel_path)
        if not records:
            print("  净值数据为空，跳过")
            continue

        data = align_data(records, benchmark_dict)
        if not data:
            continue

        # 保留既有 metrics（如果已有）
        json_path = config["json_path"]
        if json_path.exists():
            try:
                with open(json_path, 'r', encoding='utf-8') as f:
                    old_j = json.load(f)
                    if 'metrics' in old_j:
                        data['metrics'] = old_j['metrics']
            except Exception:
                pass

        os.makedirs(os.path.dirname(json_path), exist_ok=True)
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        print(f"  已保存 {json_path}")

    process_clients()
    print(f"\n{'='*50}\n更新完成！\n{'='*50}")


if __name__ == '__main__':
    main()
